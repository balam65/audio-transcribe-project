"""
Export Service
Generates transcript and summary exports in TXT, DOCX, PDF, and Markdown formats.
"""

import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

from config import config
from services.post_process import format_timestamp

logger = logging.getLogger(__name__)

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.units import inch
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


def export_transcript_txt(meeting: dict, segments: list) -> Path:
    """Export full transcript as plain text."""
    export_dir = config.get_export_dir()
    filename = _safe_filename(meeting.get("title", "meeting"), "transcript", "txt")
    filepath = export_dir / filename

    lines = [
        f"Meeting Transcript: {meeting.get('title', 'Untitled')}",
        f"Date: {meeting.get('created_at', '')}",
        f"Duration: {_format_duration(meeting.get('duration_seconds', 0))}",
        "=" * 60,
        "",
    ]

    current_speaker = ""
    for seg in segments:
        ts = format_timestamp(seg.get("start_time", 0))
        speaker = seg.get("speaker", "")
        text = seg.get("text", "")

        if speaker and speaker != current_speaker:
            lines.append("")
            lines.append(f"{speaker}:")
            current_speaker = speaker

        lines.append(f"  [{ts}] {text}")

    filepath.write_text("\n".join(lines), encoding="utf-8")
    return filepath


def export_summary_txt(meeting: dict) -> Path:
    """Export meeting summary as plain text."""
    export_dir = config.get_export_dir()
    filename = _safe_filename(meeting.get("title", "meeting"), "summary", "txt")
    filepath = export_dir / filename

    lines = [
        f"Meeting Summary: {meeting.get('title', 'Untitled')}",
        f"Date: {meeting.get('created_at', '')}",
        "=" * 60,
        "",
        "SUMMARY",
        "-" * 40,
        meeting.get("summary", "No summary available."),
        "",
        "KEY DISCUSSION POINTS",
        "-" * 40,
    ]
    for i, point in enumerate(meeting.get("key_points", []), 1):
        lines.append(f"  {i}. {point}")

    lines.extend(["", "DECISIONS MADE", "-" * 40])
    for i, dec in enumerate(meeting.get("decisions", []), 1):
        lines.append(f"  {i}. {dec}")

    lines.extend(["", "ACTION ITEMS", "-" * 40])
    for item in meeting.get("action_items", []):
        if isinstance(item, dict):
            lines.append(f"  • {item.get('task', '')} — Owner: {item.get('owner', 'Unassigned')}")
        else:
            lines.append(f"  • {item}")

    lines.extend(["", "OPEN QUESTIONS", "-" * 40])
    for q in meeting.get("open_questions", []):
        lines.append(f"  • {q}")

    filepath.write_text("\n".join(lines), encoding="utf-8")
    return filepath


def export_transcript_md(meeting: dict, segments: list) -> Path:
    """Export full transcript as Markdown."""
    export_dir = config.get_export_dir()
    filename = _safe_filename(meeting.get("title", "meeting"), "transcript", "md")
    filepath = export_dir / filename

    lines = [
        f"# Meeting Transcript: {meeting.get('title', 'Untitled')}",
        "",
        f"**Date:** {meeting.get('created_at', '')}  ",
        f"**Duration:** {_format_duration(meeting.get('duration_seconds', 0))}  ",
        "",
        "---",
        "",
    ]

    current_speaker = ""
    for seg in segments:
        ts = format_timestamp(seg.get("start_time", 0))
        speaker = seg.get("speaker", "")
        text = seg.get("text", "")

        if speaker and speaker != current_speaker:
            lines.append("")
            lines.append(f"### {speaker}")
            lines.append("")
            current_speaker = speaker

        if seg.get("is_unclear"):
            lines.append(f"> `[{ts}]` _{text}_")
        else:
            lines.append(f"> `[{ts}]` {text}")

    filepath.write_text("\n".join(lines), encoding="utf-8")
    return filepath


def export_summary_md(meeting: dict) -> Path:
    """Export meeting summary as Markdown."""
    export_dir = config.get_export_dir()
    filename = _safe_filename(meeting.get("title", "meeting"), "summary", "md")
    filepath = export_dir / filename

    lines = [
        f"# Meeting Summary: {meeting.get('title', 'Untitled')}",
        "",
        f"**Date:** {meeting.get('created_at', '')}  ",
        "",
        "---",
        "",
        "## Summary",
        "",
        meeting.get("summary", "No summary available."),
        "",
        "## Key Discussion Points",
        "",
    ]
    for point in meeting.get("key_points", []):
        lines.append(f"- {point}")

    lines.extend(["", "## Decisions Made", ""])
    for dec in meeting.get("decisions", []):
        lines.append(f"- {dec}")

    lines.extend(["", "## Action Items", ""])
    lines.append("| Task | Owner | Deadline |")
    lines.append("|------|-------|----------|")
    for item in meeting.get("action_items", []):
        if isinstance(item, dict):
            lines.append(f"| {item.get('task', '')} | {item.get('owner', 'Unassigned')} | {item.get('deadline', 'Not specified')} |")
        else:
            lines.append(f"| {item} | Unassigned | Not specified |")

    lines.extend(["", "## Open Questions", ""])
    for q in meeting.get("open_questions", []):
        lines.append(f"- {q}")

    filepath.write_text("\n".join(lines), encoding="utf-8")
    return filepath


def export_transcript_docx(meeting: dict, segments: list) -> Optional[Path]:
    """Export full transcript as DOCX."""
    if not DOCX_AVAILABLE:
        logger.warning("python-docx not installed")
        return None

    export_dir = config.get_export_dir()
    filename = _safe_filename(meeting.get("title", "meeting"), "transcript", "docx")
    filepath = export_dir / filename

    doc = Document()

    # Title
    title = doc.add_heading(f"Meeting Transcript: {meeting.get('title', 'Untitled')}", level=1)
    doc.add_paragraph(f"Date: {meeting.get('created_at', '')}")
    doc.add_paragraph(f"Duration: {_format_duration(meeting.get('duration_seconds', 0))}")
    doc.add_paragraph("─" * 50)

    current_speaker = ""
    for seg in segments:
        ts = format_timestamp(seg.get("start_time", 0))
        speaker = seg.get("speaker", "")
        text = seg.get("text", "")

        if speaker and speaker != current_speaker:
            doc.add_heading(speaker, level=2)
            current_speaker = speaker

        p = doc.add_paragraph()
        run_ts = p.add_run(f"[{ts}] ")
        run_ts.font.color.rgb = RGBColor(100, 100, 100)
        run_ts.font.size = Pt(9)

        run_text = p.add_run(text)
        if seg.get("is_unclear"):
            run_text.italic = True
            run_text.font.color.rgb = RGBColor(180, 120, 0)

    doc.save(str(filepath))
    return filepath


def export_summary_docx(meeting: dict) -> Optional[Path]:
    """Export meeting summary as DOCX."""
    if not DOCX_AVAILABLE:
        return None

    export_dir = config.get_export_dir()
    filename = _safe_filename(meeting.get("title", "meeting"), "summary", "docx")
    filepath = export_dir / filename

    doc = Document()
    doc.add_heading(f"Meeting Summary: {meeting.get('title', 'Untitled')}", level=1)
    doc.add_paragraph(f"Date: {meeting.get('created_at', '')}")
    doc.add_paragraph("─" * 50)

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(meeting.get("summary", "No summary available."))

    doc.add_heading("Key Discussion Points", level=2)
    for point in meeting.get("key_points", []):
        doc.add_paragraph(point, style="List Bullet")

    doc.add_heading("Decisions Made", level=2)
    for dec in meeting.get("decisions", []):
        doc.add_paragraph(dec, style="List Bullet")

    doc.add_heading("Action Items", level=2)
    for item in meeting.get("action_items", []):
        if isinstance(item, dict):
            doc.add_paragraph(f"{item.get('task', '')} — Owner: {item.get('owner', 'Unassigned')}", style="List Bullet")
        else:
            doc.add_paragraph(str(item), style="List Bullet")

    doc.add_heading("Open Questions", level=2)
    for q in meeting.get("open_questions", []):
        doc.add_paragraph(q, style="List Bullet")

    doc.save(str(filepath))
    return filepath


def export_transcript_pdf(meeting: dict, segments: list) -> Optional[Path]:
    """Export full transcript as PDF."""
    if not PDF_AVAILABLE:
        logger.warning("reportlab not installed")
        return None

    export_dir = config.get_export_dir()
    filename = _safe_filename(meeting.get("title", "meeting"), "transcript", "pdf")
    filepath = export_dir / filename

    doc = SimpleDocTemplate(str(filepath), pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle("CustomTitle", parent=styles["Heading1"], fontSize=16, textColor=HexColor("#1a1a2e"))
    story.append(Paragraph(f"Meeting Transcript: {meeting.get('title', 'Untitled')}", title_style))
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"Date: {meeting.get('created_at', '')}", styles["Normal"]))
    story.append(Spacer(1, 24))

    current_speaker = ""
    for seg in segments:
        ts = format_timestamp(seg.get("start_time", 0))
        speaker = seg.get("speaker", "")
        text = seg.get("text", "")

        if speaker and speaker != current_speaker:
            story.append(Spacer(1, 12))
            story.append(Paragraph(speaker, styles["Heading2"]))
            current_speaker = speaker

        color = "#b87800" if seg.get("is_unclear") else "#333333"
        para = Paragraph(
            f'<font color="#888888" size="8">[{ts}]</font> <font color="{color}">{text}</font>',
            styles["Normal"],
        )
        story.append(para)
        story.append(Spacer(1, 4))

    doc.build(story)
    return filepath


def export_summary_pdf(meeting: dict) -> Optional[Path]:
    """Export meeting summary as PDF."""
    if not PDF_AVAILABLE:
        return None

    export_dir = config.get_export_dir()
    filename = _safe_filename(meeting.get("title", "meeting"), "summary", "pdf")
    filepath = export_dir / filename

    doc = SimpleDocTemplate(str(filepath), pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(f"Meeting Summary: {meeting.get('title', 'Untitled')}", styles["Heading1"]))
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"Date: {meeting.get('created_at', '')}", styles["Normal"]))
    story.append(Spacer(1, 24))

    story.append(Paragraph("Summary", styles["Heading2"]))
    story.append(Paragraph(meeting.get("summary", "No summary."), styles["Normal"]))
    story.append(Spacer(1, 16))

    story.append(Paragraph("Key Discussion Points", styles["Heading2"]))
    for p in meeting.get("key_points", []):
        story.append(Paragraph(f"• {p}", styles["Normal"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Decisions Made", styles["Heading2"]))
    for d in meeting.get("decisions", []):
        story.append(Paragraph(f"• {d}", styles["Normal"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Action Items", styles["Heading2"]))
    for item in meeting.get("action_items", []):
        if isinstance(item, dict):
            story.append(Paragraph(f"• {item.get('task', '')} — {item.get('owner', 'Unassigned')}", styles["Normal"]))
        else:
            story.append(Paragraph(f"• {item}", styles["Normal"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Open Questions", styles["Heading2"]))
    for q in meeting.get("open_questions", []):
        story.append(Paragraph(f"• {q}", styles["Normal"]))

    doc.build(story)
    return filepath


def _safe_filename(title: str, doc_type: str, ext: str) -> str:
    """Generate a safe filename from the meeting title."""
    safe = "".join(c if c.isalnum() or c in " -_" else "" for c in title)
    safe = safe.strip().replace(" ", "_")[:50]
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    return f"{safe}_{doc_type}_{ts}.{ext}"


def _format_duration(seconds: float) -> str:
    """Format duration in seconds to human-readable string."""
    if not seconds:
        return "N/A"
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    if hours > 0:
        return f"{hours}h {minutes}m {secs}s"
    return f"{minutes}m {secs}s"
